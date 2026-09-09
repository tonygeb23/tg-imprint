"""Word documents to editor HTML, through python-docx (DECISIONS.md
decision 3, basic scope).

    read(path) -> (body_html, meta, warnings)

Read: Title and Heading 1 to 6 by style name or outline level, paragraphs,
bold, italic, underline, strike, superscript and subscript runs, the
Strong and Emphasis character styles, hyperlinks, bullet and numbered
lists from numPr with nesting from ilvl, pictures with their Word
description as the alt text (through docfile.embed_image, so they are
downscaled on the way in), tables with the first row as column headers,
and the core properties into meta.

Not read, and warned about when present: tracked changes (accept them in
Word first), footnotes and endnotes, comments, headers and footers, text
boxes, equations, and tables of contents. The body html returned here is
raw; docfile.load runs it through htmlclean.normalise.

Nothing here touches wx or prints.
"""

import html
import re

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from docx.oxml.ns import qn
from docx.text.hyperlink import Hyperlink
from docx.text.paragraph import Paragraph
from docx.table import Table

from . import docfile

NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "v": "urn:schemas-microsoft-com:vml",
    "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
}
_HEADING_RE = re.compile(r"^heading\s*([1-9])$", re.IGNORECASE)
_HEADING_ID_RE = re.compile(r"^heading([1-9])$", re.IGNORECASE)
_LIST_STYLE_RE = re.compile(r"^list (?:bullet|number|continue)\s+([2-9])$", re.IGNORECASE)

# docs/STRINGS.md, Worker A.
MSG_NOT_WORD = "That file is not a Word document TG Imprint can open. It needs a .docx file."
MSG_TRACKED = ("The document has tracked changes, which are not read here. Accept the "
               "changes in Word first, then open it again.")
MSG_FOOTNOTES = "Footnotes and endnotes were not read. Put anything you need from them in the text."
MSG_COMMENTS = "Comments were not read."
MSG_HEADERS = "Headers and footers were not read. Put anything you need from them in the text."
MSG_TEXTBOXES = "Text boxes were not read. Put their text in the body of the document."
MSG_EQUATIONS = "Equations were not read. Write them out in words, or insert them as pictures with a description."
MSG_TOC = "The table of contents was not read. The PDF gets bookmarks from the headings instead."


def _esc(text):
    return html.escape(text, quote=False)


def _style_name(paragraph):
    try:
        return paragraph.style.name or ""
    except Exception:
        return ""


def _style_id(paragraph):
    try:
        return paragraph.style.style_id or ""
    except Exception:
        return ""


def _outline_level(paragraph):
    """The outline level from the paragraph or its style chain, or None."""
    try:
        node = paragraph._p.pPr
        if node is not None:
            level = node.find(qn("w:outlineLvl"))
            if level is not None:
                return int(level.get(qn("w:val")))
    except Exception:
        pass
    try:
        style = paragraph.style
        seen = 0
        while style is not None and seen < 12:
            seen += 1
            ppr = style.element.pPr
            if ppr is not None:
                level = ppr.find(qn("w:outlineLvl"))
                if level is not None:
                    return int(level.get(qn("w:val")))
            style = style.base_style
    except Exception:
        pass
    return None


def _heading_level(paragraph):
    name = _style_name(paragraph)
    if name.lower() == "title":
        return 1
    match = _HEADING_RE.match(name) or _HEADING_ID_RE.match(_style_id(paragraph))
    if match:
        return min(int(match.group(1)), 6)
    outline = _outline_level(paragraph)
    if outline is not None and 0 <= outline <= 5:
        return outline + 1
    return None


def _num_pr(paragraph):
    """(numId, ilvl) from the paragraph or its style chain, or None."""
    try:
        ppr = paragraph._p.pPr
        if ppr is not None and ppr.numPr is not None:
            num_id = ppr.numPr.numId.val if ppr.numPr.numId is not None else None
            ilvl = ppr.numPr.ilvl.val if ppr.numPr.ilvl is not None else 0
            if num_id is not None:
                return int(num_id), int(ilvl)
    except Exception:
        pass
    try:
        style = paragraph.style
        seen = 0
        while style is not None and seen < 12:
            seen += 1
            ppr = style.element.pPr
            if ppr is not None and ppr.numPr is not None and ppr.numPr.numId is not None:
                ilvl = ppr.numPr.ilvl.val if ppr.numPr.ilvl is not None else None
                if ilvl is None:
                    # Word's own "List Bullet 2" and "List Number 3" styles
                    # are each a separate numbering at level 0; the name
                    # carries the level.
                    match = _LIST_STYLE_RE.match(_style_name(paragraph))
                    ilvl = int(match.group(1)) - 1 if match else 0
                return int(ppr.numPr.numId.val), int(ilvl)
            style = style.base_style
    except Exception:
        pass
    return None


class _Numbering:
    """numId and level to "ul" or "ol"."""

    def __init__(self, document):
        self.element = None
        try:
            self.element = document.part.numbering_part.element
        except Exception:
            self.element = None
        self._cache = {}

    def kind(self, num_id, ilvl):
        key = (num_id, ilvl)
        if key in self._cache:
            return self._cache[key]
        result = self._lookup(num_id, ilvl)
        self._cache[key] = result
        return result

    def _lookup(self, num_id, ilvl):
        if self.element is None or num_id == 0:
            return "ul" if num_id == 0 else "ul"
        try:
            num = self.element.num_having_numId(num_id)
        except Exception:
            num = None
        if num is None:
            return "ul"
        try:
            for override in num.findall(qn("w:lvlOverride")):
                if override.get(qn("w:ilvl")) == str(ilvl):
                    lvl = override.find(qn("w:lvl"))
                    fmt = lvl.find(qn("w:numFmt")) if lvl is not None else None
                    if fmt is not None:
                        return "ul" if fmt.get(qn("w:val")) == "bullet" else "ol"
            abstract_id = str(num.abstractNumId.val)
            for abstract in self.element.findall(qn("w:abstractNum")):
                if abstract.get(qn("w:abstractNumId")) != abstract_id:
                    continue
                for lvl in abstract.findall(qn("w:lvl")):
                    if lvl.get(qn("w:ilvl")) == str(ilvl):
                        fmt = lvl.find(qn("w:numFmt"))
                        value = fmt.get(qn("w:val")) if fmt is not None else "decimal"
                        return "ul" if value in ("bullet", "none") else "ol"
        except Exception:
            pass
        return "ul"


def _run_html(run):
    """One run as inline HTML."""
    text = run.text
    if not text:
        return ""
    text = _esc(text).replace("\t", " ")
    text = text.replace("\n", "<br>")
    font = run.font
    bold = run.bold
    italic = run.italic
    try:
        style_name = (run.style.name or "").lower() if run.style is not None else ""
    except Exception:
        style_name = ""
    if style_name == "strong":
        bold = True
    if style_name == "emphasis":
        italic = True
    if font.superscript:
        text = "<sup>%s</sup>" % text
    elif font.subscript:
        text = "<sub>%s</sub>" % text
    if font.strike:
        text = "<s>%s</s>" % text
    if run.underline:
        text = "<u>%s</u>" % text
    if italic:
        text = "<em>%s</em>" % text
    if bold:
        text = "<strong>%s</strong>" % text
    return text


def _pictures_in(run_element, document, state):
    """Figures for every picture in a run element, embedded."""
    out = []
    blips = run_element.findall(".//" + "{%s}blip" % NS["a"])
    imagedata = run_element.findall(".//" + "{%s}imagedata" % NS["v"])
    for node in blips + imagedata:
        rid = node.get("{%s}embed" % NS["r"]) or node.get("{%s}id" % NS["r"])
        if not rid:
            continue
        try:
            part = document.part.related_parts[rid]
            blob = part.blob
        except Exception:
            continue
        description = ""
        drawing = node
        while drawing is not None and drawing.tag != "{%s}drawing" % NS["w"]:
            drawing = drawing.getparent()
        if drawing is not None:
            doc_pr = drawing.find(".//{%s}docPr" % NS["wp"])
            if doc_pr is not None:
                description = (doc_pr.get("descr") or "").strip() or (doc_pr.get("title") or "").strip()
        try:
            embedded = docfile.embed_image(blob)
        except ValueError:
            continue
        if description:
            out.append('<figure><img src="%s" alt="%s"></figure>'
                       % (embedded.data_uri, html.escape(description, quote=True)))
        else:
            out.append('<figure><img src="%s" alt="" data-needs-alt="1"></figure>' % embedded.data_uri)
    return out


def _paragraph_inline(paragraph, document, state):
    """(inline html, [figure html]) for a paragraph's content."""
    parts = []
    figures = []
    for item in paragraph.iter_inner_content():
        if isinstance(item, Hyperlink):
            inner = "".join(_run_html(run) for run in item.runs)
            for run in item.runs:
                figures.extend(_pictures_in(run._r, document, state))
            address = (item.address or "").strip()
            if address and inner.strip():
                parts.append('<a href="%s">%s</a>' % (html.escape(address, quote=True), inner))
            else:
                parts.append(inner)
        else:
            parts.append(_run_html(item))
            figures.extend(_pictures_in(item._r, document, state))
    return "".join(parts), figures


def _table_html(table, document, state):
    rows = []
    for row_index, row in enumerate(table.rows):
        cells = []
        seen = set()
        for cell in row.cells:
            key = id(cell._tc)
            if key in seen:
                continue
            seen.add(key)
            paragraphs = []
            for paragraph in cell.paragraphs:
                inline, figures = _paragraph_inline(paragraph, document, state)
                if inline.strip():
                    paragraphs.append(inline)
                paragraphs.extend(figures)
            content = "<br>".join(paragraphs)
            span = ""
            try:
                grid = cell._tc.tcPr.gridSpan if cell._tc.tcPr is not None else None
                if grid is not None and int(grid.val) > 1:
                    span = ' colspan="%d"' % int(grid.val)
            except Exception:
                span = ""
            tag = "th" if row_index == 0 else "td"
            scope = ' scope="col"' if tag == "th" else ""
            cells.append("<%s%s%s>%s</%s>" % (tag, scope, span, content, tag))
        if cells:
            rows.append("<tr>%s</tr>" % "".join(cells))
    return "<table>%s</table>" % "".join(rows) if rows else ""


def _nest_lists(items):
    """items: [(level, tag, html)] in order. Nested lists as HTML."""
    out = []
    stack = []   # (level, tag)

    def close_to(level):
        while stack and stack[-1][0] > level:
            _lvl, tag = stack.pop()
            out.append("</li></%s>" % tag)

    for level, tag, content in items:
        if stack and level > stack[-1][0]:
            # Deeper: open a list inside the current item.
            stack.append((level, tag))
            out.append("<%s><li>%s" % (tag, content))
            continue
        close_to(level)
        if stack and stack[-1][0] == level:
            if stack[-1][1] != tag:
                _lvl, old = stack.pop()
                out.append("</li></%s>" % old)
                stack.append((level, tag))
                out.append("<%s><li>%s" % (tag, content))
            else:
                out.append("</li><li>%s" % content)
        else:
            stack.append((level, tag))
            out.append("<%s><li>%s" % (tag, content))
    close_to(-1)
    return "".join(out)


def _language(document):
    try:
        lang = (document.core_properties.language or "").strip()
        if lang:
            return lang
    except Exception:
        pass
    try:
        styles = document.styles.element
        node = styles.find(".//{%s}docDefaults//{%s}lang" % (NS["w"], NS["w"]))
        if node is not None:
            return (node.get(qn("w:val")) or "").strip()
    except Exception:
        pass
    return ""


def _scan_for_unread(document, warnings):
    body = document.element.body
    w = NS["w"]
    if body.find(".//{%s}ins" % w) is not None or body.find(".//{%s}del" % w) is not None:
        warnings.append(MSG_TRACKED)
    if body.find(".//{%s}txbxContent" % w) is not None:
        warnings.append(MSG_TEXTBOXES)
    if body.find(".//{%s}oMath" % NS["m"]) is not None:
        warnings.append(MSG_EQUATIONS)
    for instr in body.iter("{%s}instrText" % w):
        if (instr.text or "").strip().upper().startswith("TOC"):
            warnings.append(MSG_TOC)
            break
    if (body.find(".//{%s}footnoteReference" % w) is not None
            or body.find(".//{%s}endnoteReference" % w) is not None):
        warnings.append(MSG_FOOTNOTES)
    if body.find(".//{%s}commentReference" % w) is not None:
        warnings.append(MSG_COMMENTS)
    try:
        for section in document.sections:
            for part in (section.header, section.footer):
                if part is not None and any(p.text.strip() for p in part.paragraphs):
                    warnings.append(MSG_HEADERS)
                    raise StopIteration
    except StopIteration:
        pass
    except Exception:
        pass


def read(path):
    """(body_html, meta, warnings). Raises ValueError with a sentence when
    the file is not a Word document."""
    try:
        document = Document(path)
    except (PackageNotFoundError, KeyError, ValueError, OSError):
        raise ValueError(MSG_NOT_WORD)
    except Exception:
        raise ValueError(MSG_NOT_WORD)
    warnings = []
    state = {}
    numbering = _Numbering(document)
    blocks = []
    list_items = []
    title_text = ""

    def flush_list():
        if list_items:
            blocks.append(_nest_lists(list_items))
            del list_items[:]

    for item in document.iter_inner_content():
        if isinstance(item, Table):
            flush_list()
            table_html = _table_html(item, document, state)
            if table_html:
                blocks.append(table_html)
            continue
        if not isinstance(item, Paragraph):
            continue
        inline, figures = _paragraph_inline(item, document, state)
        level = _heading_level(item)
        num = _num_pr(item) if level is None else None
        if num is not None and inline.strip():
            num_id, ilvl = num
            list_items.append((ilvl, numbering.kind(num_id, ilvl), inline + "".join(figures)))
            continue
        flush_list()
        if level is not None and inline.strip():
            if _style_name(item).lower() == "title" and not title_text:
                title_text = re.sub(r"<[^>]+>", "", inline)
                title_text = html.unescape(re.sub(r"\s+", " ", title_text)).strip()
            blocks.append("<h%d>%s</h%d>" % (level, inline, level))
        elif inline.strip():
            align = ""
            try:
                from docx.enum.text import WD_ALIGN_PARAGRAPH
                value = item.alignment
                if value == WD_ALIGN_PARAGRAPH.CENTER:
                    align = ' class="align-center"'
                elif value == WD_ALIGN_PARAGRAPH.RIGHT:
                    align = ' class="align-right"'
                elif value == WD_ALIGN_PARAGRAPH.JUSTIFY:
                    align = ' class="align-justify"'
            except Exception:
                align = ""
            blocks.append("<p%s>%s</p>" % (align, inline))
        blocks.extend(figures)
    flush_list()
    _scan_for_unread(document, warnings)

    core = document.core_properties
    meta = {
        "title": (core.title or "").strip() or title_text,
        "author": (core.author or "").strip(),
        "subject": (core.subject or "").strip(),
        "lang": _language(document),
        "source_kind": "docx",
    }
    return "\n".join(blocks), meta, warnings
